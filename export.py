"""Builds an XLSX of every scored candidate from a search -- not just the
top N posted into the Slack message text (capped for message-length
reasons, see search_runner.MAX_RESULTS_PER_MESSAGE). This file gets
uploaded to Slack alongside the results message so the hiring team always
has the complete list, not just what fit in one message.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PLATFORM_LABELS = {"linkedin": "LinkedIn", "twitter": "X / Twitter", "instagram": "Instagram"}
FILTER_ICONS = {True: "PASS", False: "FAIL", None: "?"}

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="1155CC", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = ["Name", "Platform", "Profile URL", "Fit Score", "Why (Claude's rationale)", "Binary Filters", "Outreach Status", "Notes"]
COLUMN_WIDTHS = {1: 32, 2: 13, 3: 30, 4: 10, 5: 55, 6: 30, 7: 16, 8: 28}


def _filter_summary(candidate):
    results = candidate.extra.get("binary_filter_results") or []
    return "; ".join(f"{FILTER_ICONS.get(f.get('passes'), '?')}: {f.get('label', '')}" for f in results)


def build_candidates_workbook(role, candidates):
    """candidates: already-scored list (Candidate objects with .extra
    populated by scoring.score_batch). Returns the .xlsx file as bytes."""
    ranked = sorted(candidates, key=lambda c: c.extra.get("fit_score", 0), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{role.name} Candidates"[:31]  # Excel sheet-name limit

    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    for i, cand in enumerate(ranked, start=2):
        ws.cell(row=i, column=1, value=cand.name or cand.handle)
        ws.cell(row=i, column=2, value=PLATFORM_LABELS.get(cand.platform, cand.platform))
        url_cell = ws.cell(row=i, column=3, value=cand.profile_url)
        if cand.profile_url:
            url_cell.hyperlink = cand.profile_url
            url_cell.font = LINK_FONT
        ws.cell(row=i, column=4, value=cand.extra.get("fit_score", 0))
        ws.cell(row=i, column=5, value=cand.extra.get("rationale", ""))
        ws.cell(row=i, column=6, value=_filter_summary(cand))
        ws.cell(row=i, column=7, value="")
        ws.cell(row=i, column=8, value="")
        for col in (1, 2, 4, 5, 6, 7, 8):
            cell = ws.cell(row=i, column=col)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        ws.cell(row=i, column=3).border = BORDER

    for col, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.auto_filter.ref = f"A1:H{len(ranked) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
